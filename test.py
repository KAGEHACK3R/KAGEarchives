"""
KAGEarchives
Développé par : GUY KOUAKOU
"""

import os
import sqlite3
import csv
import json
import shutil
import logging
import threading
import time
from datetime import datetime, timedelta
import xml.etree.ElementTree as ET
from typing import Dict, List, Tuple

import bcrypt

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    import matplotlib.pyplot as plt
    MATPLOTLIB_AVAILABLE = True
except ImportError:
    MATPLOTLIB_AVAILABLE = False

try:
    import openpyxl
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

from flask import Flask, jsonify, request
import dearpygui.dearpygui as dpg

#########################################
# Configuration du logging
#########################################
logging.basicConfig(
    filename='kagearchives.log',
    level=logging.DEBUG,
    format='%(asctime)s - %(levelname)s - %(message)s'
)
logging.info("KAGEarchives démarré par GUY KOUAKOU")

#########################################
# Serveur API REST avec Flask
#########################################
app_api = Flask(__name__)
db_manager_global = None

@app_api.route("/archives", methods=["GET"])
def get_archives():
    archives = db_manager_global.get_all_archives()
    keys = ["id", "titre", "auteur", "reference", "date_production", "date_entree", "nature",
            "type_document", "format_fichier", "mots_cles", "langue", "confidentialite",
            "personnes_autorisees", "historique", "emplacement_physique", "chemin_numerique",
            "duree_conservation", "date_destruction", "statut", "documents_lies", "dossier_parent", "notes"]
    data = [dict(zip(keys, archive)) for archive in archives]
    return jsonify(data)

@app_api.route("/archive", methods=["POST"])
def add_archive_api():
    data = request.json
    archive_tuple = tuple(data.get(field, "") for field in [
        "titre", "auteur", "reference", "date_production", "date_entree", "nature",
        "type_document", "format_fichier", "mots_cles", "langue", "confidentialite",
        "personnes_autorisees", "historique", "emplacement_physique", "chemin_numerique",
        "duree_conservation", "date_destruction", "statut", "documents_lies", "dossier_parent", "notes"
    ])
    if db_manager_global.insert_archive(archive_tuple):
        return jsonify({"message": "Archive ajoutée"}), 201
    else:
        return jsonify({"message": "Erreur lors de l'ajout de l'archive"}), 400

def run_api():
    app_api.run(port=5000)

#########################################
# Traductions Multilingues
#########################################
TRANSLATIONS = {
    "fr": {
        "welcome": "Bienvenue dans KAGEarchives",
        "register": "Inscription",
        "login": "Connexion",
        "username": "Nom d'utilisateur",
        "password": "Mot de passe",
        "confirm_password": "Confirmer le mot de passe",
        "role": "Rôle",
        "sign_up": "S'inscrire",
        "sign_in": "Se connecter",
        "already_registered": "Déjà inscrit ? Se connecter",
        "register_prompt": "Veuillez vous inscrire",
        "login_prompt": "Veuillez vous connecter",
        "add_archive": "Ajouter l'archive",
        "update_archive": "Modifier l'archive",
        "delete_archive": "Supprimer l'archive",
        "title": "Titre",
        "author": "Auteur/Producteur",
        "reference": "Référence/Code",
        "date_production": "Date de production",
        "date_entry": "Date d'entrée",
        "nature": "Nature",
        "type_document": "Type de document",
        "file_format": "Format de fichier",
        "keywords": "Mots-clés",
        "language": "Langue",
        "confidentiality": "Confidentialité (public, privé, confidentiel)",
        "authorized_persons": "Personnes autorisées",
        "history": "Historique (modifications manuelles)",
        "physical_location": "Emplacement physique",
        "digital_path": "Chemin d'accès numérique",
        "retention_period": "Durée de conservation",
        "destruction_date": "Date de destruction ou transfert",
        "status": "Statut",
        "linked_documents": "Documents liés / Pièces jointes",
        "parent_folder": "Numéro de dossier parent",
        "notes": "Notes et observations",
        "theme": "Thème",
        "font_size": "Taille de police",
        "save_settings": "Enregistrer paramètres",
        "export_csv": "Exporter CSV",
        "export_json": "Exporter JSON",
        "export_pdf": "Exporter PDF",
        "export_xml": "Exporter XML",
        "export_excel": "Exporter Excel",
        "search": "Recherche simple",
        "reset": "Réinitialiser",
        "no_archive_selected": "Sélectionnez une archive.",
        "import_csv": "Importer CSV",
        "attach_file": "Attacher fichier",
        "clear_form": "Effacer formulaire",
        "duplicate_archive": "Dupliquer Archive",
        "preview_archive": "Aperçu Archive",
        "advanced_search": "Recherche avancée",
        "generate_report": "Générer Rapport",
        "view_history": "Voir Historique détaillé",
    },
    "en": {
        "welcome": "Welcome to KAGEarchives",
        "register": "Registration",
        "login": "Login",
        "username": "Username",
        "password": "Password",
        "confirm_password": "Confirm Password",
        "role": "Role",
        "sign_up": "Sign Up",
        "sign_in": "Sign In",
        "already_registered": "Already registered? Sign in",
        "register_prompt": "Please register",
        "login_prompt": "Please log in",
        "add_archive": "Add Archive",
        "update_archive": "Update Archive",
        "delete_archive": "Delete Archive",
        "title": "Title",
        "author": "Author/Producer",
        "reference": "Reference/Code",
        "date_production": "Production Date",
        "date_entry": "Entry Date",
        "nature": "Nature",
        "type_document": "Document Type",
        "file_format": "File Format",
        "keywords": "Keywords",
        "language": "Language",
        "confidentiality": "Confidentiality (public, private, confidential)",
        "authorized_persons": "Authorized Persons",
        "history": "History (manual modifications)",
        "physical_location": "Physical Location",
        "digital_path": "Digital Path",
        "retention_period": "Retention Period",
        "destruction_date": "Destruction or Transfer Date",
        "status": "Status",
        "linked_documents": "Linked Documents / Attachments",
        "parent_folder": "Parent Folder Number",
        "notes": "Notes and Observations",
        "theme": "Theme",
        "font_size": "Font Size",
        "save_settings": "Save Settings",
        "export_csv": "Export CSV",
        "export_json": "Export JSON",
        "export_pdf": "Export PDF",
        "export_xml": "Export XML",
        "export_excel": "Export Excel",
        "search": "Simple Search",
        "reset": "Reset",
        "no_archive_selected": "Select an archive.",
        "import_csv": "Import CSV",
        "attach_file": "Attach File",
        "clear_form": "Clear Form",
        "duplicate_archive": "Duplicate Archive",
        "preview_archive": "Preview Archive",
        "advanced_search": "Advanced Search",
        "generate_report": "Generate Report",
        "view_history": "View Detailed History",
    }
}

#########################################
# Classe de gestion de la base de données
#########################################
class DatabaseManager:
    def __init__(self, db_file="archives.db"):
        self.db_file = db_file
        self.init_db()

    def init_db(self):
        conn = sqlite3.connect(self.db_file)
        cursor = conn.cursor()
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS utilisateurs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT UNIQUE,
                password TEXT,
                role TEXT DEFAULT 'Editor'
            )
        ''')
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS archives (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                titre TEXT,
                auteur TEXT,
                reference TEXT UNIQUE,
                date_production TEXT,
                date_entree TEXT,
                nature TEXT,
                type_document TEXT,
                format_fichier TEXT,
                mots_cles TEXT,
                langue TEXT,
                confidentialite TEXT,
                personnes_autorisees TEXT,
                historique TEXT,
                emplacement_physique TEXT,
                chemin_numerique TEXT,
                duree_conservation TEXT,
                date_destruction TEXT,
                statut TEXT,
                documents_lies TEXT,
                dossier_parent TEXT,
                notes TEXT
            )
        ''')
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS historique_modifications (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                archive_id INTEGER,
                timestamp TEXT,
                user TEXT,
                description TEXT
            )
        ''')
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS audit_log (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                timestamp TEXT,
                user TEXT,
                action TEXT,
                details TEXT
            )
        ''')
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS archive_versions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                archive_id INTEGER,
                version INTEGER,
                timestamp TEXT,
                data TEXT
            )
        ''')
        conn.commit()
        conn.close()

    def hash_password(self, password: str) -> bytes:
        return bcrypt.hashpw(password.encode(), bcrypt.gensalt())

    def verify_password(self, password: str, hashed: bytes) -> bool:
        return bcrypt.checkpw(password.encode(), hashed)

    def register_user(self, username: str, password: str, confirm_password: str, role: str) -> Tuple[bool, str]:
        if password != confirm_password:
            return False, "Passwords do not match."
        hashed_pwd = self.hash_password(password)
        conn = sqlite3.connect(self.db_file)
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM utilisateurs WHERE username = ?", (username,))
        if cursor.fetchone() is not None:
            conn.close()
            return False, "Username already taken."
        cursor.execute("INSERT INTO utilisateurs (username, password, role) VALUES (?, ?, ?)",
                       (username, hashed_pwd, role))
        conn.commit()
        conn.close()
        return True, "Registration successful."

    def authenticate_user(self, username: str, password: str) -> Tuple[bool, str, str]:
        conn = sqlite3.connect(self.db_file)
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM utilisateurs WHERE username = ?", (username,))
        user = cursor.fetchone()
        conn.close()
        if user is None:
            return False, "User not found.", None
        stored_hash = user[2]
        if self.verify_password(password, stored_hash):
            return True, "Authentication successful.", user[3]
        else:
            return False, "Incorrect password.", None

    def insert_archive(self, data: tuple) -> bool:
        conn = sqlite3.connect(self.db_file)
        cursor = conn.cursor()
        try:
            cursor.execute('''
                INSERT INTO archives (
                    titre, auteur, reference, date_production, date_entree, nature,
                    type_document, format_fichier, mots_cles, langue, confidentialite,
                    personnes_autorisees, historique, emplacement_physique, chemin_numerique,
                    duree_conservation, date_destruction, statut, documents_lies, dossier_parent,
                    notes
                ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            ''', data)
            conn.commit()
            return True
        except sqlite3.IntegrityError as e:
            logging.error("Erreur insertion archive : %s", e)
            return False
        finally:
            conn.close()

    def update_archive(self, archive_id: int, data: tuple) -> bool:
        current = self.get_archive_by_id(archive_id)
        if current:
            self.save_archive_version(archive_id, current)
        conn = sqlite3.connect(self.db_file)
        cursor = conn.cursor()
        try:
            cursor.execute('''
                UPDATE archives SET
                    titre = ?, auteur = ?, reference = ?, date_production = ?, date_entree = ?,
                    nature = ?, type_document = ?, format_fichier = ?, mots_cles = ?, langue = ?,
                    confidentialite = ?, personnes_autorisees = ?, historique = ?, emplacement_physique = ?,
                    chemin_numerique = ?, duree_conservation = ?, date_destruction = ?, statut = ?,
                    documents_lies = ?, dossier_parent = ?, notes = ?
                WHERE id = ?
            ''', data + (archive_id,))
            conn.commit()
            return True
        except sqlite3.IntegrityError as e:
            logging.error("Erreur mise à jour archive : %s", e)
            return False
        finally:
            conn.close()

    def get_archive_by_id(self, archive_id: int):
        conn = sqlite3.connect(self.db_file)
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM archives WHERE id = ?", (archive_id,))
        archive = cursor.fetchone()
        conn.close()
        return archive

    def delete_archive(self, archive_id: int):
        conn = sqlite3.connect(self.db_file)
        cursor = conn.cursor()
        cursor.execute("DELETE FROM archives WHERE id = ?", (archive_id,))
        conn.commit()
        conn.close()

    def save_archive_version(self, archive_id: int, archive_data: tuple):
        conn = sqlite3.connect(self.db_file)
        cursor = conn.cursor()
        cursor.execute("SELECT COUNT(*) FROM archive_versions WHERE archive_id = ?", (archive_id,))
        count = cursor.fetchone()[0]
        version = count + 1
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        data_json = json.dumps(archive_data)
        cursor.execute('''
            INSERT INTO archive_versions (archive_id, version, timestamp, data)
            VALUES (?, ?, ?, ?)
        ''', (archive_id, version, timestamp, data_json))
        conn.commit()
        conn.close()

    def get_all_archives(self, limit: int = 50, offset: int = 0) -> List[tuple]:
        conn = sqlite3.connect(self.db_file)
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM archives LIMIT ? OFFSET ?", (limit, offset))
        results = cursor.fetchall()
        conn.close()
        return results

    def search_archives(self, keyword: str, adv_filters: dict = None) -> List[tuple]:
        conn = sqlite3.connect(self.db_file)
        cursor = conn.cursor()
        query = f"%{keyword}%"
        sql = '''
            SELECT * FROM archives WHERE 
            (titre LIKE ? OR auteur LIKE ? OR reference LIKE ? OR nature LIKE ?)
        '''
        params = [query, query, query, query]
        if adv_filters:
            if adv_filters.get("type_document"):
                sql += " AND type_document = ?"
                params.append(adv_filters["type_document"])
            if adv_filters.get("statut"):
                sql += " AND statut = ?"
                params.append(adv_filters["statut"])
            if adv_filters.get("date_entree_min"):
                sql += " AND date_entree >= ?"
                params.append(adv_filters["date_entree_min"])
            if adv_filters.get("date_entree_max"):
                sql += " AND date_entree <= ?"
                params.append(adv_filters["date_entree_max"])
        cursor.execute(sql, tuple(params))
        results = cursor.fetchall()
        conn.close()
        return results

    def log_audit(self, user: str, action: str, details: str):
        conn = sqlite3.connect(self.db_file)
        cursor = conn.cursor()
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        cursor.execute('''
            INSERT INTO audit_log (timestamp, user, action, details)
            VALUES (?, ?, ?, ?)
        ''', (timestamp, user, action, details))
        conn.commit()
        conn.close()

    def backup_db(self, backup_file: str) -> bool:
        try:
            dest_dir = os.path.dirname(backup_file)
            if dest_dir and not os.path.exists(dest_dir):
                os.makedirs(dest_dir)
            shutil.copy(self.db_file, backup_file)
            return True
        except Exception as e:
            logging.error("Erreur sauvegarde DB : %s", e)
            return False

    def export_to_csv(self, file_path: str, archive_ids: List[int] = None):
        conn = sqlite3.connect(self.db_file)
        cursor = conn.cursor()
        if archive_ids:
            placeholders = ','.join('?' for _ in archive_ids)
            cursor.execute(f"SELECT * FROM archives WHERE id IN ({placeholders})", archive_ids)
        else:
            cursor.execute("SELECT * FROM archives")
        archives = cursor.fetchall()
        conn.close()
        with open(file_path, mode="w", newline='', encoding='utf-8') as csvfile:
            writer = csv.writer(csvfile)
            header = ["id", "titre", "auteur", "reference", "date_production", "date_entree", "nature",
                      "type_document", "format_fichier", "mots_cles", "langue", "confidentialite",
                      "personnes_autorisees", "historique", "emplacement_physique", "chemin_numerique",
                      "duree_conservation", "date_destruction", "statut", "documents_lies", "dossier_parent", "notes"]
            writer.writerow(header)
            for archive in archives:
                writer.writerow(archive)

    def export_to_json(self, file_path: str):
        archives = self.get_all_archives()
        keys = ["id", "titre", "auteur", "reference", "date_production", "date_entree", "nature",
                "type_document", "format_fichier", "mots_cles", "langue", "confidentialite",
                "personnes_autorisees", "historique", "emplacement_physique", "chemin_numerique",
                "duree_conservation", "date_destruction", "statut", "documents_lies", "dossier_parent", "notes"]
        data = [dict(zip(keys, archive)) for archive in archives]
        with open(file_path, mode="w", encoding='utf-8') as jsonfile:
            json.dump(data, jsonfile, indent=4)

    def export_to_pdf(self, file_path: str):
        if not PDF_AVAILABLE:
            raise ImportError("ReportLab n'est pas installé.")
        archives = self.get_all_archives()
        c = canvas.Canvas(file_path, pagesize=letter)
        width, height = letter
        y = height - 50
        c.setFont("Helvetica", 10)
        c.drawString(50, y, "Liste des archives")
        y -= 20
        header = ["ID", "Titre", "Auteur", "Référence", "Date d'entrée"]
        c.drawString(50, y, " | ".join(header))
        y -= 15
        for archive in archives:
            line = f"{archive[0]} | {archive[1]} | {archive[2]} | {archive[3]} | {archive[5]}"
            c.drawString(50, y, line)
            y -= 15
            if y < 50:
                c.showPage()
                y = height - 50
        c.save()

    def export_to_xml(self, file_path: str):
        archives = self.get_all_archives()
        root = ET.Element("archives")
        keys = ["id", "titre", "auteur", "reference", "date_production", "date_entree", "nature",
                "type_document", "format_fichier", "mots_cles", "langue", "confidentialite",
                "personnes_autorisees", "historique", "emplacement_physique", "chemin_numerique",
                "duree_conservation", "date_destruction", "statut", "documents_lies", "dossier_parent", "notes"]
        for archive in archives:
            arch_elem = ET.SubElement(root, "archive")
            for key, value in zip(keys, archive):
                child = ET.SubElement(arch_elem, key)
                child.text = str(value)
        tree = ET.ElementTree(root)
        tree.write(file_path, encoding='utf-8', xml_declaration=True)

    def export_to_excel(self, file_path: str):
        if not EXCEL_AVAILABLE:
            raise ImportError("openpyxl n'est pas installé.")
        archives = self.get_all_archives()
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        header = ["id", "titre", "auteur", "reference", "date_production", "date_entree", "nature",
                  "type_document", "format_fichier", "mots_cles", "langue", "confidentialite",
                  "personnes_autorisees", "historique", "emplacement_physique", "chemin_numerique",
                  "duree_conservation", "date_destruction", "statut", "documents_lies", "dossier_parent", "notes"]
        ws.append(header)
        for archive in archives:
            ws.append(list(archive))
        wb.save(file_path)

#########################################
# Application et Interface Graphique
#########################################
class ArchiveApp:
    def __init__(self, db_manager: DatabaseManager):
        self.db = db_manager
        self.current_user = None
        self.current_role = None
        self.language = "fr"
        self.selected_archive_id = None
        self.font_size = 1.0
        self.theme = "Clair"
        self.selected_archives = set()
        self.current_page = 0
        self.items_per_page = 50
        self.translations = TRANSLATIONS
        self.build_ui()

    def translate(self, key: str) -> str:
        return self.translations.get(self.language, {}).get(key, key)

    def build_ui(self):
        dpg.create_context()
        self.build_menu_bar()
        self.build_register_window()
        self.build_login_window()
        self.build_main_window()
        with dpg.file_dialog(directory_selector=False, show=False, tag="export_file_dialog", callback=self.export_file_dialog_callback):
            dpg.add_file_extension(".csv")
        with dpg.file_dialog(directory_selector=False, show=False, tag="export_json_dialog", callback=self.export_json_dialog_callback):
            dpg.add_file_extension(".json")
        if PDF_AVAILABLE:
            with dpg.file_dialog(directory_selector=False, show=False, tag="export_pdf_dialog", callback=self.export_pdf_dialog_callback):
                dpg.add_file_extension(".pdf")
        with dpg.file_dialog(directory_selector=False, show=False, tag="export_xml_dialog", callback=self.export_xml_dialog_callback):
            dpg.add_file_extension(".xml")
        if EXCEL_AVAILABLE:
            with dpg.file_dialog(directory_selector=False, show=False, tag="export_excel_dialog", callback=self.export_excel_dialog_callback):
                dpg.add_file_extension(".xlsx")
        with dpg.file_dialog(directory_selector=False, show=False, tag="import_file_dialog", callback=self.import_file_dialog_callback):
            dpg.add_file_extension(".csv")
        with dpg.file_dialog(directory_selector=False, show=False, tag="attach_file_dialog", callback=self.attach_file_callback):
            dpg.add_file_extension("*.*")
        with dpg.file_dialog(directory_selector=False, show=False, tag="backup_file_dialog", callback=self.backup_file_dialog_callback):
            dpg.add_file_extension(".db")
        dpg.create_viewport(title="KAGEarchives", width=1300, height=900)
        dpg.setup_dearpygui()
        dpg.show_viewport()

    def build_menu_bar(self):
        with dpg.window(label="KAGEarchives", tag="menu_window", pos=[0, 0], width=1300, height=30):
            with dpg.menu_bar():
                with dpg.menu(label="Fichier"):
                    dpg.add_menu_item(label="Sauvegarder DB", callback=lambda s,a,u: dpg.show_item("backup_file_dialog"))
                    dpg.add_menu_item(label="Quitter", callback=lambda s,a,u: dpg.stop_dearpygui())
                with dpg.menu(label="Edition"):
                    dpg.add_menu_item(label=self.translate("duplicate_archive"), callback=self.duplicate_archive_callback)
                    dpg.add_menu_item(label=self.translate("preview_archive"), callback=self.preview_archive_callback)
                with dpg.menu(label="Aide"):
                    dpg.add_menu_item(label="À propos", callback=lambda s,a,u: dpg.show_item("about_window"))
        with dpg.window(label="À propos", tag="about_window", modal=True, show=False, width=400, height=200):
            dpg.add_text("KAGEarchives\nDéveloppé par GUY KOUAKOU")
            dpg.add_button(label="Fermer", callback=lambda s,a,u: dpg.hide_item("about_window"))

    def build_register_window(self):
        with dpg.window(label=self.translate("register"), tag="register_window", width=400, height=450, pos=[100, 100]):
            dpg.add_text(self.translate("register_prompt"))
            dpg.add_input_text(label=self.translate("username"), tag="register_username")
            dpg.add_input_text(label=self.translate("password"), tag="register_password", password=True)
            dpg.add_input_text(label=self.translate("confirm_password"), tag="register_confirm_password", password=True)
            dpg.add_combo(("Admin", "Editor", "Viewer"), label=self.translate("role"), tag="register_role", default_value="Editor")
            dpg.add_button(label=self.translate("sign_up"), callback=self.register_callback)
            dpg.add_text("", tag="register_status")
            dpg.add_button(label=self.translate("already_registered"), callback=self.show_login_from_register)

    def build_login_window(self):
        with dpg.window(label=self.translate("login"), tag="login_window", width=400, height=300, pos=[100, 100], show=False):
            dpg.add_text(self.translate("login_prompt"))
            dpg.add_input_text(label=self.translate("username"), tag="login_username")
            dpg.add_input_text(label=self.translate("password"), tag="login_password", password=True)
            dpg.add_button(label=self.translate("sign_in"), callback=self.login_callback)
            dpg.add_text("", tag="login_status")
            dpg.add_button(label=self.translate("sign_up"), callback=self.show_register_from_login)

    def show_login_from_register(self, sender, app_data, user_data):
        dpg.hide_item("register_window")
        dpg.show_item("login_window")

    def show_register_from_login(self, sender, app_data, user_data):
        dpg.hide_item("login_window")
        dpg.show_item("register_window")

    def register_callback(self, sender, app_data, user_data):
        username = dpg.get_value("register_username")
        password = dpg.get_value("register_password")
        confirm_password = dpg.get_value("register_confirm_password")
        role = dpg.get_value("register_role")
        success, message = self.db.register_user(username, password, confirm_password, role)
        dpg.set_value("register_status", message)
        self.db.log_audit(username, "Registration", message)
        if success:
            dpg.hide_item("register_window")
            dpg.show_item("login_window")

    def login_callback(self, sender, app_data, user_data):
        username = dpg.get_value("login_username")
        password = dpg.get_value("login_password")
        success, message, role = self.db.authenticate_user(username, password)
        dpg.set_value("login_status", message)
        self.db.log_audit(username, "Login", message)
        if success:
            self.current_user = username
            self.current_role = role
            dpg.hide_item("login_window")
            dpg.show_item("main_window")
            self.restrict_ui_by_role()
            self.refresh_archive_table()
            self.check_notifications()

    def restrict_ui_by_role(self):
        if self.current_role == "Viewer":
            dpg.configure_item("add_button", enabled=False)
            dpg.configure_item("update_button", enabled=False)
            dpg.configure_item("delete_button", enabled=False)
        elif self.current_role == "Editor":
            dpg.configure_item("delete_button", enabled=False)

    def build_main_window(self):
        with dpg.window(label=self.translate("welcome"), tag="main_window", width=1300, height=900, pos=[50, 50], show=False):
            dpg.add_text(self.translate("welcome"))
            with dpg.group(horizontal=True):
                dpg.add_button(label="Toggle Theme", callback=self.theme_toggle_callback)
                dpg.add_combo(("fr", "en"), label=self.translate("language"), tag="language_selector", default_value="fr", callback=self.language_selection_callback)
            with dpg.tab_bar():
                with dpg.tab(label="Informations Générales"):
                    dpg.add_input_text(label=self.translate("title"), tag="titre")
                    dpg.add_input_text(label=self.translate("author"), tag="auteur")
                    dpg.add_input_text(label=self.translate("reference"), tag="reference")
                    dpg.add_input_text(label=self.translate("date_production"), tag="date_production")
                    dpg.add_input_text(label=self.translate("date_entry"), tag="date_entree")
                    dpg.add_input_text(label=self.translate("nature"), tag="nature")
                with dpg.tab(label="Description & Métadonnées"):
                    dpg.add_input_text(label=self.translate("type_document"), tag="type_document")
                    dpg.add_input_text(label=self.translate("file_format"), tag="format_fichier")
                    dpg.add_input_text(label=self.translate("keywords"), tag="mots_cles")
                    dpg.add_input_text(label=self.translate("language"), tag="langue")
                with dpg.tab(label="Accès & Droits"):
                    dpg.add_input_text(label=self.translate("confidentiality"), tag="confidentialite")
                    dpg.add_input_text(label=self.translate("authorized_persons"), tag="personnes_autorisees")
                    dpg.add_input_text(label=self.translate("history"), tag="historique")
                with dpg.tab(label="Location Physique & Numérique"):
                    dpg.add_input_text(label=self.translate("physical_location"), tag="emplacement_physique")
                    dpg.add_input_text(label=self.translate("digital_path"), tag="chemin_numerique")
                with dpg.tab(label="Cycle de Vie & Conservation"):
                    dpg.add_input_text(label=self.translate("retention_period"), tag="duree_conservation")
                    dpg.add_input_text(label=self.translate("destruction_date"), tag="date_destruction")
                    dpg.add_input_text(label=self.translate("status"), tag="statut")
                with dpg.tab(label="Liens, Pièces jointes & Références"):
                    dpg.add_input_text(label=self.translate("linked_documents"), tag="documents_lies")
                    dpg.add_input_text(label=self.translate("parent_folder"), tag="dossier_parent")
                with dpg.tab(label="Annotations & Historique détaillé"):
                    dpg.add_input_text(label=self.translate("notes"), tag="notes", multiline=True, height=100)
                    dpg.add_button(label=self.translate("view_history"), callback=self.show_detailed_history)
                with dpg.tab(label=self.translate("advanced_search")):
                    dpg.add_input_text(label="Mot-clé", tag="adv_keyword")
                    dpg.add_input_text(label=self.translate("type_document"), tag="adv_type_document")
                    dpg.add_input_text(label=self.translate("status"), tag="adv_statut")
                    dpg.add_input_text(label="Date d'entrée min (YYYY-MM-DD)", tag="adv_date_min")
                    dpg.add_input_text(label="Date d'entrée max (YYYY-MM-DD)", tag="adv_date_max")
                    dpg.add_button(label=self.translate("advanced_search"), callback=self.advanced_search_callback)
                with dpg.tab(label="Reporting"):
                    dpg.add_button(label=self.translate("generate_report"), callback=self.generate_report)
                with dpg.tab(label="Paramètres"):
                    dpg.add_slider_float(label=self.translate("font_size"), tag="font_slider", default_value=1.0, min_value=0.8, max_value=2.0, callback=self.update_font_size)
                    dpg.add_combo(("Clair", "Sombre"), label=self.translate("theme"), tag="theme_selector", default_value="Clair", callback=self.update_theme)
                    dpg.add_button(label=self.translate("save_settings"), callback=self.save_settings)
            dpg.add_separator()
            with dpg.group(horizontal=True):
                dpg.add_button(label=self.translate("add_archive"), tag="add_button", callback=self.add_archive_callback)
                dpg.add_button(label=self.translate("update_archive"), tag="update_button", callback=self.update_archive_callback)
                dpg.add_button(label=self.translate("delete_archive"), tag="delete_button", callback=self.delete_archive_callback)
                dpg.add_button(label=self.translate("duplicate_archive"), callback=self.duplicate_archive_callback)
                dpg.add_button(label=self.translate("preview_archive"), callback=self.preview_archive_callback)
                dpg.add_button(label=self.translate("attach_file"), callback=self.open_attach_file_dialog)
                dpg.add_button(label=self.translate("import_csv"), callback=self.import_callback)
                dpg.add_button(label=self.translate("export_csv"), callback=self.export_callback)
                dpg.add_button(label=self.translate("export_json"), callback=self.export_json_callback)
                if PDF_AVAILABLE:
                    dpg.add_button(label=self.translate("export_pdf"), callback=self.export_pdf_callback)
                dpg.add_button(label=self.translate("export_xml"), callback=self.export_xml_dialog_callback)
                if EXCEL_AVAILABLE:
                    dpg.add_button(label=self.translate("export_excel"), callback=self.export_excel_dialog_callback)
                dpg.add_button(label=self.translate("clear_form"), callback=self.clear_archive_fields)
            dpg.add_text("", tag="archive_status")
            dpg.add_separator()
            with dpg.group(horizontal=True):
                dpg.add_input_text(label=self.translate("search"), tag="search_keyword", callback=self.live_search_callback)
                dpg.add_button(label=self.translate("reset"), callback=self.reset_search_callback)
            with dpg.child_window(height=250, autosize_x=True):
                with dpg.table(header_row=True, tag="archive_table", resizable=True):
                    dpg.add_table_column(label="Select", width_fixed=True)
                    dpg.add_table_column(label="ID")
                    dpg.add_table_column(label="Title")
                    dpg.add_table_column(label="Author")
                    dpg.add_table_column(label="Reference")
                    dpg.add_table_column(label="Entry Date")
            with dpg.group(horizontal=True):
                dpg.add_button(label="Previous", callback=self.prev_page)
                dpg.add_button(label="Next", callback=self.next_page)
                dpg.add_text(tag="page_info")
            with dpg.window(label="", tag="status_bar", pos=[0,850], width=1300, height=30, no_title_bar=True):
                dpg.add_text(tag="status_text", default_value="")
            threading.Thread(target=self.periodic_refresh, daemon=True).start()

    def theme_toggle_callback(self, sender, app_data, user_data):
        if self.theme == "Clair":
            with dpg.theme(tag="dark_theme"):
                with dpg.theme_component(dpg.mvAll):
                    dpg.add_theme_color(dpg.mvThemeCol_WindowBg, (30, 30, 30, 255))
                    dpg.add_theme_color(dpg.mvThemeCol_Text, (255, 255, 255, 255))
                    dpg.add_theme_color(dpg.mvThemeCol_Button, (50, 50, 50, 255))
                    dpg.add_theme_color(dpg.mvThemeCol_FrameBg, (70, 70, 70, 255))
            dpg.bind_theme("dark_theme")
            self.theme = "Sombre"
        else:
            dpg.bind_theme(None)
            self.theme = "Clair"
        dpg.set_value("theme_selector", self.theme)
        dpg.set_value("archive_status", f"Thème changé en: {self.theme}")
        self.db.log_audit(self.current_user, "Changement Thème", f"Thème défini sur {self.theme}")
        self.update_status_bar()

    def language_selection_callback(self, sender, app_data, user_data):
        self.language = dpg.get_value("language_selector")
        dpg.destroy_context()
        self.build_ui()
        dpg.set_value("archive_status", f"Langue changée en: {self.language}")

    def add_archive_callback(self, sender, app_data, user_data):
        data = (
            dpg.get_value("titre"), dpg.get_value("auteur"), dpg.get_value("reference"),
            dpg.get_value("date_production"), dpg.get_value("date_entree"), dpg.get_value("nature"),
            dpg.get_value("type_document"), dpg.get_value("format_fichier"), dpg.get_value("mots_cles"),
            dpg.get_value("langue"), dpg.get_value("confidentialite"), dpg.get_value("personnes_autorisees"),
            dpg.get_value("historique"), dpg.get_value("emplacement_physique"), dpg.get_value("chemin_numerique"),
            dpg.get_value("duree_conservation"), dpg.get_value("date_destruction"), dpg.get_value("statut"),
            dpg.get_value("documents_lies"), dpg.get_value("dossier_parent"), dpg.get_value("notes")
        )
        if not dpg.get_value("titre") or not dpg.get_value("reference"):
            dpg.set_value("archive_status", "Title and Reference are required.")
            return
        try:
            datetime.strptime(dpg.get_value("date_entree"), "%Y-%m-%d")
        except ValueError:
            dpg.set_value("archive_status", "Invalid date format (use YYYY-MM-DD).")
            return
        if self.db.insert_archive(data):
            dpg.set_value("archive_status", "Archive added successfully.")
            self.refresh_archive_table()

    def update_archive_callback(self, sender, app_data, user_data):
        if not self.selected_archive_id:
            dpg.set_value("archive_status", self.translate("no_archive_selected"))
            return
        data = (
            dpg.get_value("titre"), dpg.get_value("auteur"), dpg.get_value("reference"),
            dpg.get_value("date_production"), dpg.get_value("date_entree"), dpg.get_value("nature"),
            dpg.get_value("type_document"), dpg.get_value("format_fichier"), dpg.get_value("mots_cles"),
            dpg.get_value("langue"), dpg.get_value("confidentialite"), dpg.get_value("personnes_autorisees"),
            dpg.get_value("historique"), dpg.get_value("emplacement_physique"), dpg.get_value("chemin_numerique"),
            dpg.get_value("duree_conservation"), dpg.get_value("date_destruction"), dpg.get_value("statut"),
            dpg.get_value("documents_lies"), dpg.get_value("dossier_parent"), dpg.get_value("notes")
        )
        if self.db.update_archive(self.selected_archive_id, data):
            dpg.set_value("archive_status", "Archive updated successfully.")
            self.refresh_archive_table()

    def delete_archive_callback(self, sender, app_data, user_data):
        if not self.selected_archive_id:
            dpg.set_value("archive_status", self.translate("no_archive_selected"))
            return
        self.db.delete_archive(self.selected_archive_id)
        dpg.set_value("archive_status", "Archive deleted.")
        self.refresh_archive_table()

    def duplicate_archive_callback(self, sender, app_data, user_data):
        if self.selected_archive_id is None:
            dpg.set_value("archive_status", self.translate("no_archive_selected"))
            return
        archive = self.db.get_archive_by_id(self.selected_archive_id)
        if archive:
            new_ref = archive[3] + "_dup"
            data = (
                archive[1], archive[2], new_ref, archive[4], archive[5], archive[6],
                archive[7], archive[8], archive[9], archive[10], archive[11], archive[12],
                archive[13], archive[14], archive[15], archive[16], archive[17], archive[18],
                archive[19], archive[20], archive[21]
            )
            if self.db.insert_archive(data):
                dpg.set_value("archive_status", "Archive duplicated successfully.")
                self.db.log_audit(self.current_user, "Duplicate Archive", f"Archive {self.selected_archive_id} duplicated.")
                self.refresh_archive_table()
            else:
                dpg.set_value("archive_status", "Error duplicating archive (reference conflict).")

    def preview_archive_callback(self, sender, app_data, user_data):
        if self.selected_archive_id is None:
            dpg.set_value("archive_status", self.translate("no_archive_selected"))
            return
        archive = self.db.get_archive_by_id(self.selected_archive_id)
        if not archive:
            dpg.set_value("archive_status", "Archive not found.")
            return
        preview_text = "\n".join(f"{key}: {value}" for key, value in zip(
            ["ID", "Titre", "Auteur", "Référence", "Date Production", "Date Entrée", "Nature",
             "Type Document", "Format Fichier", "Mots-Clés", "Langue", "Confidentialité",
             "Personnes autorisées", "Historique", "Emplacement", "Chemin", "Durée", "Date Destruction",
             "Statut", "Documents liés", "Dossier Parent", "Notes"],
            archive))
        with dpg.window(label="Aperçu de l'archive", modal=True, tag="preview_window", width=600, height=400):
            dpg.add_text(preview_text)
            dpg.add_button(label="Fermer", callback=lambda s,a,u: dpg.delete_item("preview_window"))

    def view_archive_versions(self, sender, app_data, user_data):
        if not self.selected_archive_id:
            dpg.set_value("archive_status", self.translate("no_archive_selected"))
            return
        conn = sqlite3.connect(self.db.db_file)
        cursor = conn.cursor()
        cursor.execute("SELECT version, timestamp, data FROM archive_versions WHERE archive_id = ?", (self.selected_archive_id,))
        versions = cursor.fetchall()
        conn.close()
        with dpg.window(label="Archive Versions", modal=True, width=600, height=300):
            for ver in versions:
                dpg.add_text(f"Version {ver[0]} - {ver[1]}")
                dpg.add_button(label="Restore", callback=lambda s,a,u,v=ver: self.restore_version(v))

    def restore_version(self, version_data):
        archive_id = self.selected_archive_id
        data = tuple(json.loads(version_data[2])[1:])  # Exclude ID
        self.db.update_archive(archive_id, data)
        dpg.set_value("archive_status", f"Restored version {version_data[0]}")
        self.refresh_archive_table()

    def show_detailed_history(self, sender, app_data, user_data):
        if self.selected_archive_id is None:
            dpg.set_value("archive_status", "Sélectionnez une archive pour l'historique détaillé.")
            return
        conn = sqlite3.connect(self.db.db_file)
        cursor = conn.cursor()
        cursor.execute("SELECT timestamp, user, description FROM historique_modifications WHERE archive_id = ?", (self.selected_archive_id,))
        records = cursor.fetchall()
        conn.close()
        with dpg.window(label="Historique détaillé", modal=True, tag="history_window", width=500, height=300):
            for rec in records:
                dpg.add_text(f"{rec[0]} - {rec[1]}: {rec[2]}")
            dpg.add_button(label="Fermer", callback=lambda s,a,u: dpg.delete_item("history_window"))

    def advanced_search_callback(self, sender, app_data, user_data):
        keyword = dpg.get_value("adv_keyword")
        adv_filters = {
            "type_document": dpg.get_value("adv_type_document"),
            "statut": dpg.get_value("adv_statut"),
            "date_entree_min": dpg.get_value("adv_date_min"),
            "date_entree_max": dpg.get_value("adv_date_max")
        }
        results = self.db.search_archives(keyword, adv_filters)
        self.refresh_archive_table(results)

    def generate_report(self, sender, app_data, user_data):
        if not MATPLOTLIB_AVAILABLE:
            dpg.set_value("archive_status", "Matplotlib not installed, report generation unavailable.")
            return
        archives = self.db.get_all_archives()
        counts = {}
        for arch in archives:
            date_str = arch[5]
            if date_str:
                try:
                    dt = datetime.strptime(date_str, "%Y-%m-%d")
                    key = dt.strftime("%Y-%m")
                    counts[key] = counts.get(key, 0) + 1
                except:
                    continue
        months = sorted(counts.keys())
        values = [counts[m] for m in months]
        plt.figure(figsize=(8,4))
        plt.bar(months, values, color='skyblue')
        plt.xlabel("Month")
        plt.ylabel("Number of Archives")
        plt.title("Archives per Month")
        report_file = "report.png"
        plt.savefig(report_file)
        plt.close()
        width, height, channels, data = dpg.load_image(report_file)
        with dpg.window(label="Report", modal=True, tag="report_window", width=850, height=500):
            dpg.add_image("report_texture", width=width, height=height, texture_tag="report_texture")
            dpg.set_value("report_texture", data)
            dpg.add_button(label="Close", callback=lambda s,a,u: dpg.delete_item("report_window"))

    def update_font_size(self, sender, app_data, user_data):
        self.font_size = dpg.get_value("font_slider")
        dpg.set_global_font_scale(self.font_size)
        self.update_status_bar()

    def update_theme(self, sender, app_data, user_data):
        self.theme = dpg.get_value("theme_selector")
        dpg.set_value("archive_status", f"Thème mis à jour: {self.theme}")
        self.db.log_audit(self.current_user, "Changement Thème", f"Thème défini sur {self.theme}")
        self.update_status_bar()

    def save_settings(self, sender, app_data, user_data):
        settings = {"font_size": self.font_size, "theme": self.theme, "language": self.language}
        with open("settings.json", "w") as f:
            json.dump(settings, f)
        dpg.set_value("archive_status", "Paramètres sauvegardés.")

    def check_notifications(self):
        conn = sqlite3.connect(self.db.db_file)
        cursor = conn.cursor()
        cursor.execute("SELECT id, titre, date_destruction FROM archives WHERE date_destruction IS NOT NULL")
        archives = cursor.fetchall()
        conn.close()
        today = datetime.now().date()
        for archive in archives:
            try:
                destruction_date = datetime.strptime(archive[2], "%Y-%m-%d").date()
                if destruction_date - today <= timedelta(days=7):
                    with dpg.window(label="Notification", modal=True, width=300, height=100):
                        dpg.add_text(f"Archive {archive[1]} due for destruction on {archive[2]}")
                        dpg.add_button(label="Close", callback=lambda: dpg.delete_item(dpg.last_item()))
            except ValueError:
                continue

    def export_callback(self, sender, app_data, user_data):
        dpg.show_item("export_file_dialog")

    def export_file_dialog_callback(self, sender, app_data, user_data):
        if "file_path_name" in app_data:
            self.db.export_to_csv(app_data["file_path_name"], list(self.selected_archives) if self.selected_archives else None)
            dpg.set_value("archive_status", f"Exported to {app_data['file_path_name']}")

    def export_json_callback(self, sender, app_data, user_data):
        dpg.show_item("export_json_dialog")

    def export_json_dialog_callback(self, sender, app_data, user_data):
        if "file_path_name" in app_data:
            self.db.export_to_json(app_data["file_path_name"])
            dpg.set_value("archive_status", f"Exported to {app_data['file_path_name']}")

    def export_pdf_callback(self, sender, app_data, user_data):
        dpg.show_item("export_pdf_dialog")

    def export_pdf_dialog_callback(self, sender, app_data, user_data):
        if "file_path_name" in app_data:
            self.db.export_to_pdf(app_data["file_path_name"])
            dpg.set_value("archive_status", f"Exported to {app_data['file_path_name']}")

    def export_xml_dialog_callback(self, sender, app_data, user_data):
        dpg.show_item("export_xml_dialog")

    def export_xml_callback(self, sender, app_data, user_data):
        if "file_path_name" in app_data:
            self.db.export_to_xml(app_data["file_path_name"])
            dpg.set_value("archive_status", f"Exported to {app_data['file_path_name']}")

    def export_excel_dialog_callback(self, sender, app_data, user_data):
        dpg.show_item("export_excel_dialog")

    def export_excel_callback(self, sender, app_data, user_data):
        if "file_path_name" in app_data:
            self.db.export_to_excel(app_data["file_path_name"])
            dpg.set_value("archive_status", f"Exported to {app_data['file_path_name']}")

    def import_callback(self, sender, app_data, user_data):
        dpg.show_item("import_file_dialog")

    def import_file_dialog_callback(self, sender, app_data, user_data):
        if "file_path_name" in app_data:
            dpg.set_value("archive_status", f"Imported from {app_data['file_path_name']} (not implemented yet)")

    def open_attach_file_dialog(self, sender, app_data, user_data):
        dpg.show_item("attach_file_dialog")

    def attach_file_callback(self, sender, app_data, user_data):
        if "file_path_name" in app_data:
            dpg.set_value("chemin_numerique", app_data["file_path_name"])
            dpg.set_value("archive_status", f"File attached: {app_data['file_path_name']}")

    def backup_file_dialog_callback(self, sender, app_data, user_data):
        if "file_path_name" in app_data:
            if self.db.backup_db(app_data["file_path_name"]):
                dpg.set_value("archive_status", f"Database backed up to {app_data['file_path_name']}")
            else:
                dpg.set_value("archive_status", "Backup failed.")

    def clear_archive_fields(self, sender, app_data, user_data):
        for tag in ["titre", "auteur", "reference", "date_production", "date_entree", "nature",
                    "type_document", "format_fichier", "mots_cles", "langue", "confidentialite",
                    "personnes_autorisees", "historique", "emplacement_physique", "chemin_numerique",
                    "duree_conservation", "date_destruction", "statut", "documents_lies", "dossier_parent", "notes"]:
            dpg.set_value(tag, "")

    def live_search_callback(self, sender, app_data, user_data):
        self.search_archive_callback(sender, app_data, user_data)

    def search_archive_callback(self, sender, app_data, user_data):
        keyword = dpg.get_value("search_keyword")
        results = self.db.search_archives(keyword)
        self.refresh_archive_table(results)

    def reset_search_callback(self, sender, app_data, user_data):
        dpg.set_value("search_keyword", "")
        self.refresh_archive_table()

    def refresh_archive_table(self, filtered_data: List[tuple] = None):
        if filtered_data is None:
            archives = self.db.get_all_archives(limit=self.items_per_page, offset=self.current_page * self.items_per_page)
        else:
            archives = filtered_data
        dpg.delete_item("archive_table", children_only=True)
        for archive in archives:
            with dpg.table_row(parent="archive_table"):
                dpg.add_checkbox(callback=lambda s,a,u,i=archive[0]: self.toggle_selection(i))
                dpg.add_text(str(archive[0]))
                dpg.add_text(archive[1])
                dpg.add_text(archive[2])
                dpg.add_text(archive[3])
                dpg.add_text(archive[5])
                dpg.add_button(label="Select", callback=lambda s,a,u,i=archive[0]: self.load_archive(i))
        dpg.set_value("page_info", f"Page {self.current_page + 1}")

    def toggle_selection(self, archive_id):
        if archive_id in self.selected_archives:
            self.selected_archives.remove(archive_id)
        else:
            self.selected_archives.add(archive_id)

    def prev_page(self, sender, app_data, user_data):
        if self.current_page > 0:
            self.current_page -= 1
            self.refresh_archive_table()

    def next_page(self, sender, app_data, user_data):
        self.current_page += 1
        self.refresh_archive_table()

    def periodic_refresh(self):
        while True:
            time.sleep(30)
            self.refresh_archive_table()

    def update_status_bar(self):
        status = f"Utilisateur: {self.current_user if self.current_user else 'N/A'} | Rôle: {self.current_role if self.current_role else 'N/A'} | Thème: {self.theme}"
        dpg.set_value("status_text", status)

    def load_archive(self, archive_id: int):
        self.selected_archive_id = archive_id
        archive = self.db.get_archive_by_id(archive_id)
        if archive:
            tags = ["titre", "auteur", "reference", "date_production", "date_entree", "nature",
                    "type_document", "format_fichier", "mots_cles", "langue", "confidentialite",
                    "personnes_autorisees", "historique", "emplacement_physique", "chemin_numerique",
                    "duree_conservation", "date_destruction", "statut", "documents_lies", "dossier_parent", "notes"]
            for tag, value in zip(tags, archive[1:]):
                dpg.set_value(tag, value)
            dpg.set_value("archive_status", f"Loaded archive {archive_id}")
            self.update_status_bar()

def main():
    global db_manager_global
    db_manager = DatabaseManager()
    db_manager_global = db_manager
    api_thread = threading.Thread(target=run_api, daemon=True)
    api_thread.start()
    app = ArchiveApp(db_manager)
    dpg.start_dearpygui()
    dpg.destroy_context()

if __name__ == '__main__':
    main()
